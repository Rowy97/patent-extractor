from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from pathlib import Path

def output2Doc(df_us, df_us_application, df_foreign, df_non_patent, case_no = "", output_path =None):
    # Output all tables to doc
    # Create a Word document
    doc = Document()
    set_page_margins(doc)

    # Add case number
    para = doc.add_paragraph()
    para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER  # Center the paragraph
    run = para.add_run(f"Case：Our Ref {case_no}")
    run.bold = True                                  # Bold text
    run.font.size = Pt(11)                           # Optional: set font size
    run.font.color.rgb = RGBColor(0, 0, 0)           # Black text
    run.font.name = 'Times New Roman'

    # Add each DataFrame with optional title
    add_df_to_doc(doc, df_us, "U.S.PATENTS")
    add_df_to_doc(doc, df_us_application, "U.S.PATENT APPLICATION PUBLICATIONS")
    add_df_to_doc(doc, df_foreign, "FOREIGN PATENT DOCUMENTS")
    add_df_to_doc(doc, df_non_patent, "NON-PATENT LITERATURE DOCUMENTS")

    if output_path is None:
        output_path = Path("/out")   # fallback
    else:
        output_path = Path(output_path)
    # Save the document
    output_path.mkdir(parents=True, exist_ok=True)
    out_file = output_path / f"IDS_{case_no}.docx"
    doc.save(out_file)

    return str(out_file) 

def set_page_margins(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1.25)
    section.bottom_margin = Inches(1.25)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

def set_table_border(table):
    tbl = table._element

    # Get or create <w:tblPr>
    tblPr = tbl.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)

    # Remove existing <w:tblBorders> if any
    old_borders = tblPr.find(qn('w:tblBorders'))
    if old_borders is not None:
        tblPr.remove(old_borders)

    # Create new borders
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')       # border width
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), '000000')
        tblBorders.append(border)

    tblPr.append(tblBorders)

def get_column_widths(title):
    if "FOREIGN" in title.upper():
        return [Inches(x) for x in [0.5, 0.7, 0.7, 0.7, 1.3, 1.3, 1.3]]
    elif "NON-PATENT" in title.upper():
        return [Inches(x) for x in [0.5, 6]]
    else:  # U.S. PATENTS and US APPLICATIONS
        return [Inches(x) for x in [0.5, 1, 1, 1.2, 1.4, 1.4]]

def add_df_to_doc(doc, df, title):
    df = df.reset_index(drop=True)  # ensure 0..n-1

    n_cols = len(df.columns)
    n_rows = len(df) + 2  # title row + header row + data rows
    column_widths = get_column_widths(title)

    table = doc.add_table(rows=n_rows, cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_border(table)

    # set widths
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            if idx < len(column_widths):
                cell.width = column_widths[idx]
                tc = cell._tc
                tcPr = tc.get_or_add_tcPr()
                tcW = OxmlElement('w:tcW')
                tcW.set(qn('w:type'), 'dxa')
                tcW.set(qn('w:w'), str(int(column_widths[idx].inches * 1440)))
                tcPr.append(tcW)

    # title (merged)
    cell = table.cell(0, 0)
    cell.text = title
    cell.merge(table.cell(0, n_cols - 1))
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for p in cell.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if p.runs:
            r = p.runs[0]
            r.bold = True
            r.font.size = Pt(11)
            r.font.name = 'Times New Roman'

    # headers
    hdr_cells = table.rows[1].cells
    for i, col_name in enumerate(df.columns):
        hdr_cells[i].text = str(col_name)
        for p in hdr_cells[i].paragraphs:
            if p.runs:
                r = p.runs[0]
                r.bold = True
                r.font.size = Pt(11)
                r.font.name = 'Times New Roman'

    # data rows (use enumerate to avoid index issues)
    for r, (_, row) in enumerate(df.iterrows(), start=0):
        row_cells = table.rows[r + 2].cells
        for c, value in enumerate(row):
            row_cells[c].text = "" if value is None else str(value)
            for p in row_cells[c].paragraphs:
                if p.runs:
                    rr = p.runs[0]
                    rr.font.size = Pt(11)
                    rr.font.name = 'Times New Roman'

    doc.add_paragraph()

